"""Extração de texto de arquivos baixados do Moodle."""

from __future__ import annotations

import io
import re
from email.message import Message

from bs4 import BeautifulSoup


def filename_from(resp) -> str:
    cd = resp.headers.get("content-disposition", "")
    if cd:
        msg = Message()
        msg["content-disposition"] = cd
        name = msg.get_filename()
        if name:
            return name
    return resp.url.path.rsplit("/", 1)[-1] or "arquivo"


def extract_text(content: bytes, filename: str, content_type: str = "") -> str:
    name = filename.lower()
    ctype = content_type.lower()
    if name.endswith(".pdf") or "pdf" in ctype:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        pages = [f"--- página {i + 1} ---\n{p.extract_text() or ''}" for i, p in enumerate(reader.pages)]
        return "\n".join(pages)
    if name.endswith(".docx") or "wordprocessingml" in ctype:
        import docx

        doc = docx.Document(io.BytesIO(content))
        return "\n".join(p.text for p in doc.paragraphs)
    if name.endswith(".pptx") or "presentationml" in ctype:
        from pptx import Presentation

        prs = Presentation(io.BytesIO(content))
        slides = []
        for i, slide in enumerate(prs.slides):
            textos = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            slides.append(f"--- slide {i + 1} ---\n" + "\n".join(textos))
        return "\n".join(slides)
    if name.endswith((".html", ".htm")) or "html" in ctype:
        return html_to_text(content.decode("utf-8", errors="replace"))
    if ctype.startswith("text/") or name.endswith((".txt", ".md", ".csv", ".py", ".c", ".java", ".json", ".tex")):
        return content.decode("utf-8", errors="replace")
    raise ValueError(f"Formato não suportado para leitura de texto: {filename} ({content_type}). Use baixar_arquivo.")


def html_to_text(html: str) -> str:
    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["script", "style"]):
        tag.decompose()
    text = soup.get_text("\n", strip=True)
    return re.sub(r"\n{3,}", "\n\n", text)
